import csv
import io
import os
from dataclasses import dataclass, field
from pathlib import Path

import pypdfium2 as pdfium
import pytesseract
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from pytesseract import TesseractNotFoundError


TESSERACT_CMD = os.getenv("TESSERACT_CMD")
if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD

MAX_SPREADSHEET_ROWS = 20_000
MIN_NATIVE_PDF_TEXT_CHARACTERS = 40


class DocumentExtractionError(RuntimeError):
    """Raised when a supported document cannot be converted to text."""


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    extraction_method: str
    item_count: int
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)


def _clean_text(parts: list[str]) -> str:
    return "\n\n".join(part.strip() for part in parts if part and part.strip()).strip()


def _ocr_image(image: Image.Image) -> str:
    try:
        return pytesseract.image_to_string(image, lang="eng").strip()
    except TesseractNotFoundError as error:
        raise DocumentExtractionError(
            "OCR is required for this file, but the Tesseract executable was not found. "
            "Install Tesseract and set TESSERACT_CMD in backend/.env."
        ) from error


def _extract_pdf(file_path: Path) -> ExtractionResult:
    try:
        reader = PdfReader(str(file_path))
        rendered_pdf = pdfium.PdfDocument(str(file_path))
    except Exception as error:
        raise DocumentExtractionError("The PDF could not be opened") from error

    pages: list[str] = []
    ocr_pages: list[int] = []

    try:
        for index, page in enumerate(reader.pages):
            native_text = (page.extract_text() or "").strip()

            if len(native_text) >= MIN_NATIVE_PDF_TEXT_CHARACTERS:
                pages.append(f"--- Page {index + 1} ---\n{native_text}")
                continue

            try:
                bitmap = rendered_pdf[index].render(scale=2.0)
                image = bitmap.to_pil()
                ocr_text = _ocr_image(image)
            except DocumentExtractionError:
                raise
            except Exception as error:
                raise DocumentExtractionError(
                    f"OCR failed for PDF page {index + 1}"
                ) from error

            ocr_pages.append(index + 1)
            combined = ocr_text or native_text
            pages.append(f"--- Page {index + 1} ---\n{combined}")
    finally:
        rendered_pdf.close()

    text = _clean_text(pages)
    if not text:
        raise DocumentExtractionError("No readable text was found in the PDF")

    warnings: list[str] = []
    method = "pypdf"
    if ocr_pages:
        method = "pypdf+tesseract_ocr"
        warnings.append(
            "OCR was used on PDF page(s): " + ", ".join(map(str, ocr_pages))
        )

    return ExtractionResult(
        text=text,
        extraction_method=method,
        item_count=len(reader.pages),
        page_count=len(reader.pages),
        warnings=warnings,
    )


def _extract_image(file_path: Path) -> ExtractionResult:
    try:
        with Image.open(file_path) as image:
            text = _ocr_image(image.convert("RGB"))
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError("The image could not be read") from error

    if not text:
        raise DocumentExtractionError("OCR did not find readable text in the image")

    return ExtractionResult(
        text=text,
        extraction_method="tesseract_ocr",
        item_count=1,
        page_count=1,
    )


def _extract_docx(file_path: Path) -> ExtractionResult:
    try:
        document = Document(str(file_path))
    except Exception as error:
        raise DocumentExtractionError("The DOCX file could not be opened") from error

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    text = _clean_text(parts)
    if not text:
        raise DocumentExtractionError("No readable text was found in the DOCX file")

    return ExtractionResult(
        text=text,
        extraction_method="python_docx",
        item_count=len(parts),
    )


def _extract_xlsx(file_path: Path) -> ExtractionResult:
    try:
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    except Exception as error:
        raise DocumentExtractionError("The XLSX workbook could not be opened") from error

    lines: list[str] = []
    row_count = 0
    warnings: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            lines.append(f"--- Sheet: {worksheet.title} ---")
            for row in worksheet.iter_rows(values_only=True):
                if row_count >= MAX_SPREADSHEET_ROWS:
                    warnings.append(
                        f"Extraction stopped after {MAX_SPREADSHEET_ROWS} spreadsheet rows"
                    )
                    break
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    lines.append(" | ".join(values))
                    row_count += 1
            if row_count >= MAX_SPREADSHEET_ROWS:
                break
    finally:
        workbook.close()

    text = _clean_text(lines)
    if not text:
        raise DocumentExtractionError("No populated cells were found in the XLSX workbook")

    return ExtractionResult(
        text=text,
        extraction_method="openpyxl",
        item_count=row_count,
        warnings=warnings,
    )


def _extract_csv(file_path: Path) -> ExtractionResult:
    try:
        content = file_path.read_text(encoding="utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        lines: list[str] = []
        for index, row in enumerate(reader):
            if index >= MAX_SPREADSHEET_ROWS:
                break
            if any(cell.strip() for cell in row):
                lines.append(" | ".join(cell.strip() for cell in row))
    except (OSError, UnicodeError, csv.Error) as error:
        raise DocumentExtractionError("The CSV file could not be parsed") from error

    text = _clean_text(lines)
    if not text:
        raise DocumentExtractionError("No rows were found in the CSV file")

    warnings = []
    if len(lines) >= MAX_SPREADSHEET_ROWS:
        warnings.append(f"Extraction stopped after {MAX_SPREADSHEET_ROWS} CSV rows")

    return ExtractionResult(
        text=text,
        extraction_method="python_csv",
        item_count=len(lines),
        warnings=warnings,
    )


def _extract_plain_text(file_path: Path, extension: str) -> ExtractionResult:
    try:
        text = file_path.read_text(encoding="utf-8-sig").strip()
    except (OSError, UnicodeError) as error:
        raise DocumentExtractionError("The text file could not be decoded as UTF-8") from error

    if not text:
        raise DocumentExtractionError("The text file is empty")

    method = "markdown" if extension == ".md" else "plain_text"
    return ExtractionResult(
        text=text,
        extraction_method=method,
        item_count=max(1, len(text.splitlines())),
    )


def extract_document(file_path: Path, extension: str) -> ExtractionResult:
    """Dispatch one already-validated upload to its format-specific extractor."""
    normalized_extension = extension.lower()

    if normalized_extension == ".pdf":
        return _extract_pdf(file_path)
    if normalized_extension in {".png", ".jpg", ".jpeg"}:
        return _extract_image(file_path)
    if normalized_extension == ".docx":
        return _extract_docx(file_path)
    if normalized_extension == ".xlsx":
        return _extract_xlsx(file_path)
    if normalized_extension == ".csv":
        return _extract_csv(file_path)
    if normalized_extension in {".txt", ".md"}:
        return _extract_plain_text(file_path, normalized_extension)

    raise DocumentExtractionError(f"No extractor exists for {normalized_extension}")
