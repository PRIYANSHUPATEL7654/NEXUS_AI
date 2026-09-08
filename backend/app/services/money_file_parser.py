import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from dateutil import parser as date_parser
from openpyxl import load_workbook
import pdfplumber

from .document_extraction import DocumentExtractionError, extract_document
from .transaction_parser import parse_transaction_text


SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".pdf"}
ALLOWED_CATEGORIES = {
    "Food", "Transport", "Shopping", "Bills", "Health", "Education",
    "Entertainment", "Housing", "Salary", "Transfer", "Other",
}
HEADER_ALIASES = {
    "date": {"date", "transaction date", "txn date", "value date"},
    "description": {"description", "details", "narration", "merchant", "particulars", "remarks"},
    "amount": {"amount", "transaction amount", "value"},
    "debit": {"debit", "withdrawal", "withdrawal amount", "paid"},
    "credit": {"credit", "deposit", "deposit amount", "received"},
    "type": {"type", "transaction type", "kind"},
    "category": {"category", "expense category"},
}


@dataclass(frozen=True)
class FileTransaction:
    transaction_date: date
    description: str
    amount: Decimal
    transaction_type: str
    category: str
    confidence: Decimal
    source_location: str
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class FileParseResult:
    transactions: list[FileTransaction]
    warnings: list[str]
    extraction_method: str


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _column_map(headers: list[object]) -> dict[str, int]:
    result: dict[str, int] = {}
    normalized = [_normalize_header(value) for value in headers]
    for canonical, aliases in HEADER_ALIASES.items():
        for index, header in enumerate(normalized):
            if header in aliases:
                result[canonical] = index
                break
    return result


def _cell(row: list[object], columns: dict[str, int], name: str) -> object | None:
    index = columns.get(name)
    return row[index] if index is not None and index < len(row) else None


def _decimal(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    cleaned = str(value).replace(",", "").replace("₹", "").replace("INR", "").replace("Rs.", "").strip()
    try:
        return Decimal(cleaned).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        return None
    try:
        return date_parser.parse(str(value), dayfirst=True).date()
    except (ValueError, OverflowError):
        return None


def _category(description: str, amount: Decimal, transaction_date: date, supplied: object) -> tuple[str, bool]:
    supplied_text = str(supplied or "").strip().title()
    if supplied_text in ALLOWED_CATEGORIES:
        return supplied_text, True
    sample = parse_transaction_text(
        f"{description} INR {amount} on {transaction_date.isoformat()}",
        today=transaction_date,
    ).transactions[0]
    return sample.category, sample.category != "Other"


def _parse_rows(rows: list[list[object]], source_name: str) -> tuple[list[FileTransaction], list[str]]:
    if not rows:
        return [], [f"{source_name} contained no rows"]
    columns = _column_map(rows[0])
    if "date" not in columns or "description" not in columns or not ({"amount", "debit", "credit"} & columns.keys()):
        return [], [
            f"{source_name} needs Date, Description and either Amount or Debit/Credit columns"
        ]

    transactions: list[FileTransaction] = []
    warnings: list[str] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in row):
            continue
        parsed_date = _date(_cell(row, columns, "date"))
        description = str(_cell(row, columns, "description") or "").strip()
        amount = _decimal(_cell(row, columns, "amount"))
        debit = _decimal(_cell(row, columns, "debit"))
        credit = _decimal(_cell(row, columns, "credit"))

        supplied_type = str(_cell(row, columns, "type") or "").strip().lower()
        if credit is not None and credit > 0:
            amount, transaction_type = credit, "income"
        elif debit is not None and debit > 0:
            amount, transaction_type = debit, "expense"
        elif amount is not None:
            transaction_type = supplied_type if supplied_type in {"income", "expense"} else ("expense" if amount >= 0 else "income")
            amount = abs(amount)
        else:
            warnings.append(f"{source_name}, row {row_number} was skipped: invalid amount")
            continue

        if parsed_date is None or not description or amount <= 0:
            warnings.append(f"{source_name}, row {row_number} was skipped: date, description or amount is invalid")
            continue

        category, category_recognized = _category(
            description,
            amount,
            parsed_date,
            _cell(row, columns, "category"),
        )
        item_warnings: list[str] = []
        confidence = Decimal("0.95")
        if supplied_type and supplied_type not in {"income", "expense"}:
            item_warnings.append("Unrecognized type was replaced using debit/credit or amount rules")
            confidence -= Decimal("0.10")
        if not category_recognized:
            item_warnings.append("No category matched; Other was proposed")
            confidence -= Decimal("0.15")

        transactions.append(
            FileTransaction(
                transaction_date=parsed_date,
                description=description[:500],
                amount=amount,
                transaction_type=transaction_type,
                category=category,
                confidence=confidence,
                source_location=f"{source_name}, row {row_number}",
                warnings=item_warnings,
            )
        )
    return transactions, warnings


def _parse_csv(path: Path) -> FileParseResult:
    try:
        content = path.read_text(encoding="utf-8-sig")
        rows = [list(row) for row in csv.reader(io.StringIO(content))]
    except (OSError, UnicodeError, csv.Error) as error:
        raise ValueError("The CSV file could not be parsed") from error
    transactions, warnings = _parse_rows(rows, "CSV")
    return FileParseResult(transactions, warnings, "python_csv_columns_v1")


def _parse_xlsx(path: Path) -> FileParseResult:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("The XLSX workbook could not be opened") from error
    transactions: list[FileTransaction] = []
    warnings: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            parsed, sheet_warnings = _parse_rows(rows, f"Sheet {sheet.title}")
            transactions.extend(parsed)
            warnings.extend(sheet_warnings)
    finally:
        workbook.close()
    return FileParseResult(transactions, warnings, "openpyxl_columns_v1")


def _parse_pdf(path: Path) -> FileParseResult:
    table_transactions: list[FileTransaction] = []
    table_warnings: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                for table_number, table in enumerate(page.extract_tables(), start=1):
                    rows = [[cell or "" for cell in row] for row in table if row]
                    parsed_rows, warnings = _parse_rows(
                        rows, f"PDF page {page_number}, table {table_number}"
                    )
                    table_transactions.extend(parsed_rows)
                    table_warnings.extend(warnings)
    except Exception:
        table_warnings.append("PDF table detection failed; text extraction fallback was used")

    if table_transactions:
        return FileParseResult(table_transactions, table_warnings, "pdfplumber_tables_v1")

    try:
        extraction = extract_document(path, ".pdf")
    except DocumentExtractionError as error:
        raise ValueError(str(error)) from error
    parsed = parse_transaction_text(extraction.text)
    transactions = [
        FileTransaction(
            transaction_date=item.transaction_date,
            description=item.description,
            amount=item.amount,
            transaction_type=item.transaction_type,
            category=item.category,
            confidence=item.confidence,
            source_location="PDF extracted text",
            warnings=item.warnings,
        )
        for item in parsed.transactions
    ]
    return FileParseResult(
        transactions,
        extraction.warnings + parsed.warnings,
        extraction.extraction_method + "+transaction_rules_v1",
    )


def parse_money_file(path: Path, extension: str) -> FileParseResult:
    normalized = extension.lower()
    if normalized not in SUPPORTED_EXTENSIONS:
        raise ValueError("Money file imports support PDF, XLSX and CSV")
    result = _parse_csv(path) if normalized == ".csv" else _parse_xlsx(path) if normalized == ".xlsx" else _parse_pdf(path)
    if not result.transactions:
        details = "; ".join(result.warnings) or "No recognizable transactions were found"
        raise ValueError(details)
    return result
